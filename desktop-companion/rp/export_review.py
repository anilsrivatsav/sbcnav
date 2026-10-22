"""Portable Excel review export with checked, cached formulas."""
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import os
import tempfile
import zipfile
from xml.etree import ElementTree as E
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

HEADERS = ['Auction date', 'Category', 'Lot / asset', 'Station / train', 'Rate unit', 'Proposed value', 'Justification', 'Base / working rate', 'Factor / second-bidder rate', 'Calculation status', 'Previous auction date', 'Previous outcome', 'Accepted LAR / rate', 'Awarded party', 'Contract status', 'Contract award', 'Contract start', 'Contract end', 'Foreclosed date', 'Catalogue', 'Source / missing evidence', 'Rule version']
Q = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def calculated(formula, values):
    h, i = Decimal(str(values[7])), Decimal(str(values[8]))
    operations = {
        '=+H{r}': lambda: h,
        '=+H{r}*1.1': lambda: h * Decimal('1.1'),
        '=+H{r}*0.9': lambda: h * Decimal('.9'),
        '=+H{r}*I{r}': lambda: h * i,
        '=ROUND(H{r}*I{r},2)': lambda: h * i,
        '=IF(H{r}>I{r},H{r},I{r})': lambda: max(h, i),
        '=IF(H{r}<300000,300000,H{r})': lambda: max(h, Decimal(300000)),
    }
    if formula not in operations:
        raise ValueError('The Excel review contains an unsupported formula.')
    return operations[formula]().quantize(Decimal('.01'), rounding=ROUND_HALF_UP)


def export_review(output, rows):
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Review'
    sheet.append(HEADERS)
    cached = {}
    for index, row in enumerate(rows, 2):
        values = list(row['values'])
        if len(values) != len(HEADERS):
            raise ValueError('The Excel review columns changed.')
        formula = row.get('formula')
        if formula:
            value = calculated(formula, values)
            if abs(value - Decimal(str(row['expected']))) > Decimal('.005'):
                raise ValueError('Calculation does not reconcile for ' + str(values[2]))
            cached['F' + str(index)] = str(value)
        for col, value in enumerate(values):
            if col in (0, 10, 15, 16, 17, 18) and isinstance(value, str) and len(value) == 10:
                try:
                    value = date.fromisoformat(value)
                except ValueError:
                    pass
            if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
                value = "'" + value
            values[col] = value
        sheet.append(values)
        if formula:
            sheet[f'F{index}'] = formula.replace('{r}', str(index))
    last = max(2, len(rows) + 1)
    for row in sheet.iter_rows(min_row=1, max_row=last, max_col=22):
        for cell in row:
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 75
    for cell in sheet[1]:
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='24445F')
    sheet.row_dimensions[1].height = 42
    for column in 'ABCDEFGHIJKLMNOPQRSTUV':
        sheet.column_dimensions[column].width = 20
    for column, width in [('G', 95), ('C', 45), ('U', 60)]:
        sheet.column_dimensions[column].width = width
    for index in range(2, last + 1):
        for col in ('F', 'H', 'I', 'M'):
            sheet[f'{col}{index}'].number_format = '#,##0.00'
        for col in ('A', 'K', 'P', 'Q', 'R', 'S'):
            sheet[f'{col}{index}'].number_format = 'dd-mm-yyyy'
    table = Table(displayName='DeskFlowReview', ref=f'A1:V{last}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = 'D2'
    output = Path(output)
    with tempfile.TemporaryDirectory(dir=output.parent) as folder:
        original = Path(folder) / 'review.xlsx'
        checked = Path(folder) / 'verified.xlsx'
        wb.save(original)
        wb.close()
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(checked, 'w', zipfile.ZIP_DEFLATED) as target:
            for entry in source.infolist():
                data = source.read(entry.filename)
                if entry.filename == 'xl/worksheets/sheet1.xml':
                    root = E.fromstring(data)
                    for cell in root.iter(Q + 'c'):
                        if cell.get('r') in cached:
                            value = cell.find(Q + 'v')
                            if value is None:
                                value = E.SubElement(cell, Q + 'v')
                            value.text = cached[cell.get('r')]
                    data = E.tostring(root, encoding='utf-8', xml_declaration=True)
                target.writestr(entry, data)
        verify = load_workbook(checked, data_only=True)
        try:
            for address, value in cached.items():
                if Decimal(str(verify.active[address].value)) != Decimal(value):
                    raise ValueError('The exported Excel result did not verify.')
        finally:
            verify.close()
        os.replace(checked, output)

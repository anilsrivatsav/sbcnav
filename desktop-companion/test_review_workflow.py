"""Exercise review, guarded writeback and software readback on disposable books."""
import json
import tempfile
import threading
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import Mock

import openpyxl
from deskflow import Deskflow, reader_lifetime
from rp.store import Store
from rp.workbooks import SPECS
from rp.review import digest


class ReviewWorkflowTests(unittest.TestCase):
    def test_navigation_failure_closes_only_the_owned_reader_tab(self):
        owned=Mock();owned.is_closed.return_value=False
        unrelated=Mock()
        with self.assertRaisesRegex(RuntimeError,'navigation failed'):
            with reader_lifetime() as pages:
                pages.append(owned)
                raise RuntimeError('navigation failed')
        owned.close.assert_called_once()
        unrelated.close.assert_not_called()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.paths = {}
        for group, spec in SPECS.items():
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = spec['sheet']
            for cell, value in spec['checks'].items():
                sheet[cell] = value
            path = self.root / (group + '.xlsx')
            book.save(path)
            book.close()
            self.paths[group] = path
        (self.root / 'workbooks.json').write_text(json.dumps({k:str(v) for k,v in self.paths.items()}))
        self.d = object.__new__(Deskflow)
        self.d.userdir = self.root
        self.d.store = Store(self.root / 'test.sqlite')
        self.d.lock = threading.RLock()
        self.d.thread = None
        self.d.state = {}
        self.d._workbook_digests = {}
        self.d.updater = Mock()
        self.d.updater.snapshot.return_value = {'busy':False}
        today = date.today().isoformat()
        self.key = self.d.store.save_lot({'catalogue':'TEST-C','lot':'TEST-L','description':'Fixture asset',
            'start':(date.today()+timedelta(days=5)).isoformat(),'status':'RP Pending','rate_unit':'Annual Licensing Fee'})
        self.d.store.save_evidence('__workspace__','scheduled_snapshot',{'ids':[self.key],'complete':True,'count':1,'collected_at':today})
        self.d.store.save_evidence('__workspace__','history_snapshot',{'complete':True,'from_date':'2022-01-01','to_date':today,'collected_at':today})
        scan=self.d.library().start_scan('2022-01-01',today)
        self.d.library().finish_scan(scan,{'complete':True,'expected':0,'visited':[]})

    def evidence(self):
        record = self.d.detail(self.key)
        self.d.command('save-review',{'key':self.key,'revision':record['revision'],
            'inputs':{'family':'Parking / Pay & Use','benchmark_source':'Fixture approved benchmark',
                'initial_rp':'1000','new_asset':True,'identity_verified':True,'conditions_reviewed':True},'contract':{}})
        return self.d.detail(self.key)

    def test_review_save_and_software_readback_have_identical_amount_and_note(self):
        record = self.evidence()
        self.assertTrue(record['result']['ready'])
        before = {k:digest(v) for k,v in self.paths.items()}
        report = self.d.write_workbooks(ready_only=True)
        self.assertEqual(report['saved'],[])
        self.assertEqual(before,{k:digest(v) for k,v in self.paths.items()})
        record = self.d.detail(self.key)
        self.d.command('review-proposal',{'key':self.key,'revision':record['revision'],'proposal_revision':record['proposal']['revision']})
        proposed = self.d.overview()['lots'][0]['proposal']
        report = self.d.write_workbooks(ready_only=True)
        self.assertEqual(len(report['saved']),1)
        self.assertEqual(digest(self.paths['parcel']),before['parcel'])
        self.assertEqual(digest(self.paths['catering']),before['catering'])
        book = openpyxl.load_workbook(self.paths['commercial'],data_only=True)
        self.assertEqual(book['Table1']['E6'].value,float(proposed['value']))
        self.assertEqual(book['Table1']['F6'].value,proposed['justification'])
        book.close()
        row = self.d.overview()['lots'][0]
        self.assertEqual(row['value'],float(proposed['value']))
        self.assertEqual(row['justification'],proposed['justification'])
        self.assertEqual(row['proposal']['status'],'saved')

    def test_changed_workbook_requires_another_review(self):
        record = self.evidence()
        self.d.command('review-proposal',{'key':self.key,'revision':record['revision'],'proposal_revision':record['proposal']['revision']})
        # Insert an existing manual value after review, as a concurrent Excel edit.
        book = openpyxl.load_workbook(self.paths['commercial'])
        sheet = book['Table1']
        for cell,value in {'L6':'TEST-L','N6':'TEST-C','G6':'Annual Licensing Fee','E6':900,'F6':'Manual review'}.items():sheet[cell]=value
        book.save(self.paths['commercial']);book.close()
        report = self.d.write_workbooks(ready_only=True)
        self.assertEqual(report['saved'],[])
        book = openpyxl.load_workbook(self.paths['commercial'],data_only=True)
        self.assertEqual(book['Table1']['E6'].value,900)
        self.assertEqual(book['Table1']['F6'].value,'Manual review')
        book.close()

    def test_review_rejects_forged_history_or_stale_revision(self):
        record = self.evidence()
        for patch in ({'identity':'unknown','bid_rate':'9'}, {'identity':'unknown','outcome':'Sold'}):
            with self.assertRaises(ValueError):
                self.d.command('save-review',{'key':self.key,'revision':record['revision'],'inputs':{},'contract':{},'auctions':[patch]})
        with self.assertRaisesRegex(ValueError,'Evidence changed'):
            self.d.command('save-review',{'key':self.key,'revision':'old','inputs':{},'contract':{}})

    def test_legacy_snapshot_cannot_bypass_incomplete_library(self):
        record=self.evidence()
        with self.d.store.connect() as c:c.execute('DELETE FROM df_scans')
        with self.assertRaisesRegex(ValueError,'one-time outcome backfill'):
            self.d.command('review-proposal',{'key':self.key,'revision':record['revision'],'proposal_revision':record['proposal']['revision']})


if __name__ == '__main__':
    unittest.main()

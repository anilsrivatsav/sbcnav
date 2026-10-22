import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import urllib.error
import portable
from rp.export_review import export_review
from openpyxl import load_workbook


class ConnectionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.file = self.folder / 'Connection.sbcnav'
        self.token = 'fixture-credential-not-for-production-' * 2
        self.file.write_text(json.dumps({'format': 'sbcnav-connection-v1', 'operator_token': self.token, 'url': 'https://untrusted.invalid'}))

    def test_file_cannot_redirect_credential_to_another_host(self):
        from io import BytesIO
        with patch.object(portable, 'STATE', self.folder / 'data'), patch.object(portable.urllib.request, 'urlopen', return_value=BytesIO(b'{"data":{}}')) as api:
            portable.install_connection(self.file)
        request = api.call_args.args[0]
        self.assertEqual(request.full_url, portable.API + '/session')
        self.assertEqual((self.folder / 'data/operator.key').read_text(), self.token)
        self.assertNotIn(self.token, (self.folder / 'data/connection.json').read_text())

    def test_rejected_key_keeps_previous_connection(self):
        (self.folder / 'operator.key').write_text('previous-working-credential')
        error = urllib.error.HTTPError(portable.API + '/session', 401, 'Rejected', {}, None)
        with patch.object(portable, 'STATE', self.folder), patch.object(portable.urllib.request, 'urlopen', side_effect=error):
            with self.assertRaisesRegex(ValueError, 'no longer accepted'):
                portable.install_connection(self.file)
        self.assertEqual((self.folder / 'operator.key').read_text(), 'previous-working-credential')

    def test_invalid_format_and_injected_token_are_rejected(self):
        for value in [{'format': 'wrong', 'operator_token': self.token}, {'format': 'sbcnav-connection-v1', 'operator_token': self.token + '\r\nX: y'}]:
            self.file.write_text(json.dumps(value))
            with self.assertRaises(ValueError):
                portable.connection_token(self.file)

    def test_missing_connection_at_startup_is_quiet(self):
        with patch.object(portable, 'STATE', self.folder), patch.object(portable.sys, 'argv', ['app', '--background']), patch.object(portable, 'setup') as setup, patch.object(portable, 'start_service') as start:
            portable.main()
        setup.assert_not_called()
        start.assert_not_called()


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.output = Path(self.temp.name) / 'review.xlsx'

    def row(self, formula, h, i, expected):
        return {'values': ['2026-10-22', 'Parking', '=untrusted()', '', 'year', expected, 'Evidence', h, i] + [''] * 13, 'formula': formula, 'expected': expected}

    def test_all_pricing_formulas_keep_formula_and_verified_cached_amount(self):
        rows = [self.row(*item) for item in [
            ('=+H{r}', 100, 1, 100), ('=+H{r}*1.1', 100, 1, 110),
            ('=+H{r}*0.9', 100, .9, 90), ('=+H{r}*I{r}', 100, .7, 70),
            ('=ROUND(H{r}*I{r},2)', 100.15, .9, 90.14),
            ('=IF(H{r}>I{r},H{r},I{r})', 100, 150, 150),
            ('=IF(H{r}<300000,300000,H{r})', 100, 1, 300000)]]
        export_review(self.output, rows)
        values = load_workbook(self.output, data_only=True)
        formulas = load_workbook(self.output, data_only=False)
        try:
            for index, row in enumerate(rows, 2):
                self.assertEqual(values.active[f'F{index}'].value, row['expected'])
                self.assertEqual(formulas.active[f'F{index}'].value, row['formula'].replace('{r}', str(index)))
                self.assertEqual(formulas.active[f'C{index}'].data_type, 's')
            self.assertEqual(values.active.freeze_panes, 'D2')
        finally:
            values.close()
            formulas.close()

    def test_mismatch_does_not_replace_an_existing_export(self):
        self.output.write_bytes(b'previous export')
        with self.assertRaisesRegex(ValueError, 'does not reconcile'):
            export_review(self.output, [self.row('=+H{r}*I{r}', 100, .9, 100)])
        self.assertEqual(self.output.read_bytes(), b'previous export')

    def test_arbitrary_formula_is_rejected(self):
        with self.assertRaisesRegex(ValueError, 'unsupported formula'):
            export_review(self.output, [self.row('=WEBSERVICE("https://untrusted.invalid")', 100, 1, 100)])

    def test_empty_review_remains_a_valid_excel_workbook(self):
        export_review(self.output, [])
        book = load_workbook(self.output)
        self.assertEqual(book.active['A1'].value, 'Auction date')
        book.close()


if __name__ == '__main__':
    unittest.main()

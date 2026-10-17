"""Google Sheets through the user's browser session; no Google API/OAuth app."""
from io import BytesIO, StringIO
import csv
import time
import re
import zipfile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from core import cell, literal, serial
from openpyxl.utils import get_column_letter,range_boundaries

def quote(title):
    return "'" + title.replace("'", "''") + "'"

def export(context, spreadsheet):
    url = f'https://docs.google.com/spreadsheets/d/{spreadsheet}/export?format=xlsx'
    response = context.request.get(url, timeout=120000)
    data = response.body()
    if not response.ok or not data.startswith(b'PK'):
        raise RuntimeError('Google Sheet export failed. Sign in to the Google account with edit access in the app Chrome window.')
    return data

def workbook(data):
    formula = load_workbook(BytesIO(data), data_only=False, read_only=True)
    computed = load_workbook(BytesIO(data), data_only=True, read_only=True)
    return formula, computed

def filter_spec(data,title):
    """Read the public XLSX filter definition before browser pastes."""
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(BytesIO(data)) as z:
        book=ET.fromstring(z.read('xl/workbook.xml'))
        names=[s.attrib['name'] for s in book.find('s:sheets',ns)]
        root=ET.fromstring(z.read(f'xl/worksheets/sheet{names.index(title)+1}.xml'))
        af=root.find('s:autoFilter',ns)
        if af is None:return None
        result={'ref':af.attrib['ref'].replace('$',''),'columns':[]}
        for column in af:
            values=column.find('s:filters',ns)
            if values is None or any(c.tag!=f"{{{ns['s']}}}filter" for c in values):
                raise RuntimeError('This sheet has an unsupported filter condition. Clear that filter before applying; the saved preview is kept.')
            result['columns'].append({'column':int(column.attrib['colId'])+1,
                'values':[c.attrib['val'] for c in values], 'blank':values.attrib.get('blank')=='1'})
        return result

def grid(formula, computed, title, columns):
    if title not in formula.sheetnames:
        raise RuntimeError('Expected sheet tab missing: '+title)
    f, v = formula[title], computed[title]
    result = []
    for frow, vrow in zip(f.iter_rows(max_col=columns), v.iter_rows(max_col=columns)):
        out = []
        for fc, vc in zip(frow, vrow):
            raw, effective = fc.value, vc.value
            if fc.is_date and raw is not None and not fc.data_type == 'f':
                raw = serial(raw.date() if hasattr(raw,'date') else raw)
            if vc.is_date and effective is not None:
                effective = serial(effective.date() if hasattr(effective,'date') else effective)
            c = cell(raw) if fc.data_type=='f' else literal(raw)
            if effective is not None:
                c['effectiveValue'] = literal(effective).get('userEnteredValue',{})
            if getattr(fc,'comment',None):
                c['note'] = fc.comment.text
            out.append(c)
        result.append(out)
    return result

def tsv(rows):
    def text(c):
        v = c.get('userEnteredValue', {})
        if 'formulaValue' in v:
            return v['formulaValue']
        if 'numberValue' in v:
            return str(v['numberValue'])
        s = str(v.get('stringValue',''))
        # Treat all source strings literally, including formula-like invoice descriptions.
        if s:
            s = "'"+s
        return s
    # Sheets accepts quoted TSV cells. Flattening remarks would change manual
    # notes and make the export verification fail after an otherwise valid paste.
    output = StringIO(newline='')
    writer = csv.writer(output,delimiter='\t',lineterminator='\n')
    writer.writerows([[text(c) for c in row] for row in rows])
    return output.getvalue().removesuffix('\n')

class SheetUI:
    def __init__(self, context, spreadsheet, title, gid, log):
        self.context, self.spreadsheet, self.title, self.gid, self.log = context, spreadsheet, title, gid, log
        self.page = context.new_page()
        self.page.goto(f'https://docs.google.com/spreadsheets/d/{spreadsheet}/edit#gid={gid}')
        self.page.locator('#t-name-box').wait_for(timeout=120000)
        self.page.set_default_timeout(30000)

    def assert_editable(self):
        body=self.page.locator('body').inner_text()
        if 'View only' in body or 'Request edit access' in body:
            raise RuntimeError('This Google account cannot edit '+self.title+'. Switch accounts, then Resume the saved preview.')

    def suspend_filter(self,spec):
        if not spec:return
        self.log('Temporarily showing all rows for verified updates: '+self.title)
        self.page.get_by_role('menuitem',name='Data',exact=True).click()
        self.page.get_by_role('menuitem').filter(has_text=re.compile('^Remove filter$')).click()
        self.saved()
        if filter_spec(export(self.context,self.spreadsheet),self.title) is not None:
            raise RuntimeError('The sheet filter was not removed. No filtered paste was attempted.')

    def restore_filter(self,spec):
        if not spec:return
        current=filter_spec(export(self.context,self.spreadsheet),self.title)
        if current==spec:return
        if current is not None:
            self.suspend_filter(current)
        self.select(spec['ref'])
        self.page.get_by_role('menuitem',name='Data',exact=True).click()
        self.page.get_by_role('menuitem').filter(has_text=re.compile('^Create a filter$')).click()
        self.saved()
        first_col,first_row,_,_=range_boundaries(spec['ref'])
        for condition in spec['columns']:
            self.select(f"{get_column_letter(first_col+condition['column']-1)}{first_row}")
            self.page.keyboard.press('Control+Alt+R')
            self.page.get_by_text('Clear',exact=True).click()
            for value in condition['values']+(['(Blanks)'] if condition['blank'] else []):
                self.page.get_by_text(value,exact=True).click()
            self.page.get_by_text('OK',exact=True).click()
            self.saved()
        restored=filter_spec(export(self.context,self.spreadsheet),self.title)
        if restored!=spec:
            raise RuntimeError('Values were verified, but the original filter needs restoration. Resume the saved preview.')
        self.log('Original filter restored: '+self.title)

    def select(self, a1):
        field = self.page.locator('#t-name-box')
        field.fill(a1)
        field.press('Enter')
        self.page.locator('#waffle-rich-text-editor').wait_for(state='attached')

    def paste(self, a1, rows):
        if not rows:
            return
        self.assert_editable()
        self.select(a1)
        self.context.grant_permissions(['clipboard-read','clipboard-write'], origin='https://docs.google.com')
        self.page.evaluate('(text)=>navigator.clipboard.writeText(text)', tsv(rows))
        self.page.keyboard.press('Control+V')
        # Allow Sheets to finish handling native paste; export/readback is the actual verification.
        time.sleep(0.7)
        body = self.page.locator('body').inner_text()
        if 'View only' in body or 'Request edit access' in body:
            raise RuntimeError('Google Sheet is not editable with this account.')

    def copy_format(self, source, destination):
        self.select(source)
        self.page.keyboard.press('Control+C')
        self.select(destination)
        self.page.keyboard.press('Control+Alt+V')
        time.sleep(0.4)

    def new_tab(self, title):
        self.page.get_by_role('button', name='Add Sheet', exact=True).click()
        time.sleep(0.7)
        # Find the active tab through the visible tab bar, not internal sheet APIs.
        self.rename_active_tab(title)

    def rename_active_tab(self,title):
        tab = self.page.locator('.docs-sheet-tab.docs-sheet-active-tab')
        if tab.count() != 1:
            raise RuntimeError('Could not locate the new status tab. No data was pasted into it.')
        tab.dblclick()
        field = tab.locator('.docs-sheet-tab-name[contenteditable="true"]')
        if field.count() != 1:
            field = self.page.get_by_role('textbox', name='Rename sheet')
        field.fill(title)
        field.press('Enter')
        time.sleep(0.5)
        if tab.locator('.docs-sheet-tab-name').inner_text()!=title:
            raise RuntimeError('The status tab rename did not complete. No status data was pasted.')

    def existing_tab(self, title):
        tab = self.page.locator('.docs-sheet-tab').filter(has_text=title)
        if tab.count() != 1:
            # Name box navigation works even if a visible sheet tab is off-screen.
            self.select(quote(title)+'!A1')
        else:
            tab.click()
        time.sleep(0.3)

    def saved(self):
        # Sheets enforces Trusted Types; its page context rejects the eval
        # used by Playwright's string wait_for_function polling.
        deadline=time.monotonic()+60
        while time.monotonic()<deadline:
            text=self.page.locator('body').inner_text()
            if 'Saving…' not in text and 'Saving...' not in text:
                time.sleep(1)
                return
            time.sleep(1)
        raise RuntimeError('Google Sheets is still saving. Resume after it finishes; no repeat paste was attempted.')

    def close(self):
        self.page.close()

def verify(context, spreadsheet, title, patches):
    """Verify exactly the written literals/formulas with authenticated XLSX export."""
    last_errors = []
    for attempt in range(3):
        data = export(context, spreadsheet)
        f, v = workbook(data)
        g = grid(f,v,title,37 if title == 'E Auction' else 28 if title == 'Main Sheet' else 10)
        errors = []
        for start_row, start_col, rows in patches:
            for i,row in enumerate(rows):
                for j,c in enumerate(row):
                    actual = g[start_row-1+i][start_col-1+j] if start_row-1+i < len(g) else {}
                    ev, av = c.get('userEnteredValue',{}), actual.get('userEnteredValue',{})
                    if ev != av:
                        if 'numberValue' in ev and 'numberValue' in av and abs(ev['numberValue']-av['numberValue']) < 0.001:
                            continue
                        errors.append((start_row+i,start_col+j))
        if not errors:
            return data
        last_errors = errors
        time.sleep(2)
    raise RuntimeError(f'Write verification failed in {title} at {last_errors[:8]}. No automatic repeat-paste was performed. Check the run report before retrying.')

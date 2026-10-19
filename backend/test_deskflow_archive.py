import hashlib
import io
import unittest
from datetime import datetime
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook

from deskflow_archive import _archive_source, _record_from_row, update_archive


class WorkbookArchiveTests(unittest.TestCase):
    def test_legacy_commercial_description_without_short_asset_is_retained(self):
        row = {'n': 69, 'v': ['13.03.24', 'ADVT MA INT', None, '180MEMU'],
               'f': {}, 's': '13.03.24 advt ma int 180memu'}
        record = _record_from_row({'id': 'others', 'portfolio': 'Commercial'}, '29.12.23', row)
        self.assertEqual(record['asset'], '180MEMU')
        self.assertEqual(record['date'], '2024-03-13')

    def workbook(self, title):
        book = Workbook()
        sheet = book.active
        sheet.title = title
        sheet.append(['Date', 'Category', 'Asset'])
        sheet.append([datetime(2026, 10, 19), 'Parking', 'Test asset'])
        sheet['H2'] = '=1+2'
        sheet.row_dimensions[2].hidden = True
        # Excel's formatted but empty rows must not become row-dimension objects.
        for row in range(3, 20003):
            sheet.row_dimensions[row].height = 15
        output = io.BytesIO()
        book.save(output)
        return output.getvalue()

    def test_empty_dimensions_are_removed_but_source_cells_are_unchanged(self):
        raw = self.workbook('Sheet1')
        compact = _archive_source(raw)
        with ZipFile(io.BytesIO(raw)) as before, ZipFile(io.BytesIO(compact)) as after:
            source_xml = before.read('xl/worksheets/sheet1.xml')
            compact_xml = after.read('xl/worksheets/sheet1.xml')
            self.assertGreater(source_xml.count(b'<row '), 20000)
            self.assertEqual(compact_xml.count(b'<row '), 2)
        book = load_workbook(io.BytesIO(compact))
        self.assertEqual(book.active['A2'].value, datetime(2026, 10, 19))
        self.assertEqual(book.active['H2'].value, '=1+2')
        self.assertTrue(book.active.row_dimensions[2].hidden)
        self.assertEqual(len(book.active.row_dimensions), 1)

    def test_archive_keeps_original_hash_and_indexes_excel_dates(self):
        sources = {key: self.workbook(title) for key, title in
                   [('parcels', 'Sheet1 (2)'), ('catering', 'Sheet1'), ('others', '29.12.23')]}
        previous = {
            'workbooks': [{'id': key, 'portfolio': key, 'file': key + '.xlsx', 'sheets': []}
                          for key in sources],
            'sourceManifest': {key: {'sha256': ''} for key in sources},
            'auctionRecords': [],
        }
        package = io.BytesIO()
        with ZipFile(package, 'w', ZIP_DEFLATED) as bundle:
            for key, raw in sources.items():
                bundle.writestr(key + '.xlsx', raw)
        archive, changed = update_archive(previous, package.getvalue())
        self.assertTrue(changed)
        self.assertEqual(archive['totalRows'], 6)
        self.assertEqual(archive['auctionRecordCount'], 3)
        self.assertTrue(all(row['date'] == '2026-10-19' for row in archive['auctionRecords']))
        for key, raw in sources.items():
            self.assertEqual(archive['sourceManifest'][key]['sha256'], hashlib.sha256(raw).hexdigest())
        same, changed = update_archive(archive, package.getvalue())
        self.assertFalse(changed)
        self.assertIs(same, archive)
        # A legacy extract may contain the original row without an indexed
        # auction record. A refresh must recover that record as well.
        archive['auctionRecords'] = []
        archive['sourceManifest'] = {key: {'sha256': ''} for key in sources}
        repaired, changed = update_archive(archive, package.getvalue())
        self.assertTrue(changed)
        self.assertEqual(repaired['auctionRecordCount'], 3)


if __name__ == '__main__':
    unittest.main()

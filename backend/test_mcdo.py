import os
os.environ.setdefault('DATABASE_URL','sqlite:///:memory:')
import unittest
from copy import deepcopy
from unittest.mock import patch
from datetime import date
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from fastapi import HTTPException
import mcdo_router as api
from models import McdoSyncRun
from mcdo.core import cell, entered
from mcdo.sheets_api import computed_workbook
from openpyxl import load_workbook
from io import BytesIO

class FakeSheets:
    def __init__(self, grids):self.grids=deepcopy(grids);self.calls=[];self.fail_once=True;self.status={}
    def read(self,name,title=None,columns=None):
        if title:return {'grid':deepcopy(self.status.get(name,[]))}
        return {'properties':{'sheetId':1,'gridProperties':{'rowCount':10}},'grid':deepcopy(self.grids[name])}
    def patches(self,name,current,patches,formats,status):
        if name=='master' and self.fail_once:self.fail_once=False;raise RuntimeError('Second sheet unavailable')
        self.calls.append((name,deepcopy(patches)))
        from mcdo.checkpoint import planned_grid
        self.grids[name]=planned_grid(self.grids[name],patches,api.SOURCES[name]['columns'])
        self.status[name]=deepcopy(status)

class McdoTests(unittest.TestCase):
    def setUp(self):
        self.engine=create_engine('sqlite:///:memory:')
        McdoSyncRun.__table__.create(self.engine)
        self.sessions=sessionmaker(self.engine,expire_on_commit=False)
        self.patches=[patch.object(api,'engine',self.engine),patch.object(api,'SessionLocal',self.sessions)]
        for p in self.patches:p.start()
    def tearDown(self):
        for p in reversed(self.patches):p.stop()
        self.engine.dispose()
    def run_fixture(self):
        grids={name:[[cell('old')]+[{}]*(spec['columns']-1)] for name,spec in api.SOURCES.items()}
        plan={'earnings':[(2,1,[[cell('new')]])],'master':[(1,1,[[cell('new')]])],'status':[[cell('Verified')]],'formats':[],'summary':{}}
        payload={'before':{n:{'grid':g} for n,g in grids.items()},'plan':plan,'end':api.today().strftime('%d/%m/%Y')}
        with self.sessions.begin() as session:session.add(McdoSyncRun(run_id='test',state='applying',payload=payload,result={}))
        return grids
    def state(self):
        with self.sessions() as session:return api.find_run(session,'test').state
    def test_manual_edit_aborts_both_writes(self):
        grids=self.run_fixture();fake=FakeSheets(grids);fake.grids['master'][0][0]=cell('manual edit')
        with patch.object(api,'Sheets',return_value=fake),patch.object(api,'import_sources') as importer:api.apply_worker('test')
        self.assertEqual(self.state(),'needs_review');self.assertEqual(fake.calls,[]);importer.assert_not_called()
    def test_partial_write_resumes_without_reappending(self):
        fake=FakeSheets(self.run_fixture())
        with patch.object(api,'Sheets',return_value=fake),patch.object(api,'import_sources',return_value={'verified':True}) as importer:
            api.apply_worker('test');self.assertEqual(self.state(),'needs_review');importer.assert_not_called()
            api.apply_worker('test');self.assertEqual(self.state(),'complete');importer.assert_called_once()
        self.assertEqual(entered(fake.grids['earnings'][1][0]),'new');self.assertEqual(len(fake.grids['earnings']),2)
        earnings=[p for name,p in fake.calls if name=='earnings']
        self.assertEqual(len(earnings[0]),1);self.assertEqual(earnings[1],[])
    def test_no_database_import_when_readback_fails(self):
        fake=FakeSheets(self.run_fixture());fake.fail_once=False
        original=fake.patches
        def discard(name,*args):
            original(name,*args)
            if name=='earnings':fake.grids[name][1][0]={}
        with patch.object(api,'Sheets',return_value=fake),patch.object(fake,'patches',side_effect=discard),patch.object(api,'import_sources') as importer:api.apply_worker('test')
        self.assertEqual(self.state(),'needs_review');importer.assert_not_called()
    def test_old_preview_is_refused(self):
        fake=FakeSheets(self.run_fixture())
        with self.sessions.begin() as session:
            r=api.find_run(session,'test');r.payload={**r.payload,'end':'01/01/2020'}
        with patch.object(api,'Sheets',return_value=fake):api.apply_worker('test')
        self.assertEqual(fake.calls,[]);self.assertEqual(self.state(),'needs_review')
    def test_incomplete_ireps_scan_is_refused(self):
        with self.assertRaisesRegex(ValueError,'incomplete'):api.evidence_plan({}, {'completed_categories':['Advertising']})
    def test_admin_access_is_closed_by_default(self):
        with patch.dict(os.environ,{'MCDO_ADMIN_TOKEN':''}):
            with self.assertRaises(HTTPException) as err:api.administrator(None)
            self.assertEqual(err.exception.status_code,503)
        with patch.dict(os.environ,{'MCDO_ADMIN_TOKEN':'x'*40}):
            with self.assertRaises(HTTPException):api.administrator('Bearer bad')
            api.administrator('Bearer '+'x'*40)
    def test_master_general_format_date_serials_are_preserved(self):
        row=[{} for _ in range(37)]
        row[5]={'effectiveValue':{'numberValue':46301}}
        row[17]={'effectiveValue':{'numberValue':46301}}
        raw=computed_workbook('master',{'grid':[row]})
        sheet=load_workbook(BytesIO(raw),data_only=True).active
        self.assertEqual(sheet.cell(1,6).value.date(),date(2026,10,6))
        self.assertEqual(sheet.cell(1,18).value.date(),date(2026,10,6))

if __name__=='__main__':unittest.main()

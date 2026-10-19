"""Encryption and source reconciliation for the Deskflow workbook archive."""
from __future__ import annotations

import hashlib
import io
import json
import os
import re
import zlib
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from cryptography.fernet import Fernet, InvalidToken
from openpyxl import load_workbook


def _fernet() -> Fernet:
    key = os.getenv('DESKFLOW_ARCHIVE_KEY', '')
    if not key:
        raise RuntimeError('DESKFLOW_ARCHIVE_KEY is required to read workbook history.')
    try:
        return Fernet(key.encode())
    except (ValueError, TypeError) as exc:
        raise RuntimeError('DESKFLOW_ARCHIVE_KEY is invalid.') from exc


def encrypt_archive(payload: dict) -> str:
    raw = json.dumps(payload, ensure_ascii=False, separators=(',', ':'), allow_nan=False).encode()
    return _fernet().encrypt(zlib.compress(raw, level=9)).decode()


@lru_cache(maxsize=2)
def decrypt_archive(ciphertext: str) -> dict:
    try:
        return json.loads(zlib.decompress(_fernet().decrypt(ciphertext.encode())))
    except (InvalidToken, ValueError, zlib.error) as exc:
        raise RuntimeError('Workbook archive could not be decrypted. Check DESKFLOW_ARCHIVE_KEY.') from exc


def _serialise(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float, bool)):
        return value
    return str(value)


_EMPTY_ROWS = re.compile(
    rb'<(?P<prefix>(?:[A-Za-z_][\w.-]*:)?)row\b[^>]*(?:/\s*>|>\s*</(?P=prefix)row\s*>)'
)


def _archive_source(raw: bytes) -> bytes:
    """Omit empty row formatting from the temporary copy used for extraction.

    Excel can save row dimensions through row 1,048,576. Loading those empty
    dimensions twice consumes far more memory than the populated workbook.
    Rows containing cells remain byte-for-byte intact, including hidden flags.
    The original upload is still used for the source hash and is never edited.
    """
    memory = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as source, ZipFile(memory, 'w') as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if re.fullmatch(r'xl/worksheets/sheet\d+\.xml', entry.filename):
                content = _EMPTY_ROWS.sub(b'', content)
            target.writestr(entry, content)
    return memory.getvalue()


def update_archive(previous: dict, bundle: bytes) -> tuple[dict, bool]:
    """Rebuild the encrypted archive input from the three selected Windows files."""
    expected = {book['id']: book for book in previous['workbooks']}
    if set(expected) != {'parcels', 'catering', 'others'}:
        raise ValueError('The existing workbook archive is incomplete.')
    names = {book_id: f'{book_id}.xlsx' for book_id in expected}
    try:
        with ZipFile(io.BytesIO(bundle)) as package:
            if len(package.namelist()) != len(names) or set(package.namelist()) != set(names.values()):
                raise ValueError('Upload exactly the three selected workbooks.')
            if any(package.getinfo(name).file_size > 12_000_000 for name in names.values()):
                raise ValueError('A workbook exceeds the 12 MB limit.')
            if sum(package.getinfo(name).file_size for name in names.values()) > 24_000_000:
                raise ValueError('The workbooks exceed the 24 MB expanded limit.')
            sources = {book_id: package.read(name) for book_id, name in names.items()}
    except (BadZipFile, KeyError) as exc:
        raise ValueError('Invalid workbook bundle.') from exc
    hashes = {book_id: hashlib.sha256(raw).hexdigest() for book_id, raw in sources.items()}
    if all(previous['sourceManifest'][book_id]['sha256'] == digest for book_id, digest in hashes.items()):
        return previous, False

    old_rows = {(book['id'], sheet['name'], row['n']): row
                for book in previous['workbooks'] for sheet in book['sheets'] for row in sheet['rows']}
    old_records = {(record['bookId'], record['sheet'], record['row']): record
                   for record in previous['auctionRecords']}
    books, records, manifest = [], [], {}
    for book_id, old_book in expected.items():
        raw = sources[book_id]
        try:
            extraction_source = _archive_source(raw)
            formula_book = load_workbook(io.BytesIO(extraction_source), data_only=False, read_only=False)
            value_book = load_workbook(io.BytesIO(extraction_source), data_only=True, read_only=False)
        except (ValueError, BadZipFile, KeyError) as exc:
            raise ValueError(f'Cannot read {old_book["file"]}.') from exc
        try:
            required = {'parcels': 'Sheet1 (2)', 'catering': 'Sheet1', 'others': '29.12.23'}[book_id]
            if required not in formula_book.sheetnames:
                raise ValueError(f'{old_book["file"]} is missing {required}.')
            sheets, cell_count = [], 0
            for sheet in formula_book.worksheets:
                cached_sheet = value_book[sheet.title]
                rows = []
                for n in sorted({cell.row for cell in sheet._cells.values() if cell.value is not None}):
                    values, formulas = [], {}
                    for column in range(1, sheet.max_column + 1):
                        cell = sheet._cells.get((n, column))
                        value = cell.value if cell else None
                        if value is not None:
                            cell_count += 1
                        if isinstance(value, str) and value.startswith('='):
                            formulas[str(column)] = value
                            cached = cached_sheet._cells.get((n, column))
                            values.append(_serialise(cached.value if cached else None))
                        else:
                            values.append(_serialise(value))
                    while values and values[-1] is None:
                        values.pop()
                    row = {'n': n, 'v': values, 'f': formulas,
                           'hidden': bool(sheet.row_dimensions[n].hidden),
                           's': ' '.join(str(value) for value in values if value is not None).lower()}
                    rows.append(row)
                    identity = (book_id, sheet.title, n)
                    prior = old_rows.get(identity)
                    if prior and identity in old_records and all(prior[field] == row[field] for field in ('v', 'f', 'hidden')):
                        records.append(old_records[identity])
                    elif identity in old_records or sheet.title == required:
                        record = _record_from_row(old_book, sheet.title, row, old_records.get(identity))
                        if record:
                            records.append(record)
                sheets.append({'name': sheet.title, 'rows': rows})
            book = {**old_book, 'sheets': sheets, 'rowCount': sum(len(s['rows']) for s in sheets),
                    'cellCount': cell_count}
            books.append(book)
            manifest[book_id] = {'file': book['file'], 'sha256': hashes[book_id],
                                 'sheets': len(sheets), 'rows': book['rowCount'], 'cells': cell_count}
        finally:
            formula_book.close()
            value_book.close()
    assets_by_key = {}
    for record in records:
        key = record['key']
        if key not in assets_by_key:
            assets_by_key[key] = {field: record.get(field) for field in
                ('key', 'bookId', 'portfolio', 'asset', 'unit', 'assetType')}
            assets_by_key[key].update(count=0, latestDate=record['date'])
        asset = assets_by_key[key]
        asset['count'] += 1
        asset['latestDate'] = max(asset['latestDate'], record['date'])
    assets = list(assets_by_key.values())
    for asset in assets:
        asset['search'] = ' '.join(str(asset[field] or '') for field in
            ('portfolio', 'asset', 'unit', 'assetType')).lower()
    updated = {**previous, 'workbooks': books, 'auctionRecords': records, 'assets': assets,
               'sourceManifest': manifest, 'generatedAt': datetime.now().astimezone().isoformat(),
               'totalRows': sum(book['rowCount'] for book in books),
               'totalCells': sum(book['cellCount'] for book in books),
               'auctionRecordCount': len(records)}
    return updated, True


def _record_from_row(book: dict, sheet: str, row: dict, prior: dict | None = None) -> dict | None:
    values = row['v']
    def at(column):
        return values[column - 1] if len(values) >= column else None
    date_label = str(at(1) or '').strip()
    if not date_label:
        return None
    parsed = None
    for fmt in ('%Y-%m-%d', '%d.%m.%Y', '%d.%m.%y', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            parsed = datetime.strptime(date_label.split('T')[0].split(' ')[0], fmt).date()
            break
        except ValueError:
            pass
    if parsed is None:
        return None
    book_id = book['id']
    config = {'parcels': (3, 13, 16), 'catering': (3, 12, 13), 'others': (3, 8, 9)}[book_id]
    source_asset = str(at(config[0]) or '').strip()
    if book_id == 'others' and not source_asset:
        # Older commercial rows identify the asset only in the description
        # column. Keep those records when worksheet insertions move their rows.
        source_asset = str(at(4) or '').strip()
    asset = source_asset
    if not asset:
        return None
    rate = at(config[1])
    try:
        rate = float(str(rate).replace(',', '')) if rate not in (None, '') else None
    except (TypeError, ValueError):
        rate = None
    if book_id == 'parcels':
        kind = 'VP' if re.search(r'\bVP\d*\b', asset, re.I) else 'FSLR' if re.search(r'\b(?:SLR|FSLR)\b', asset, re.I) else 'Parcel'
        position = 'Rear' if re.search(r'\b(?:rear|rslr)\b', asset, re.I) else 'Front' if kind == 'FSLR' else ''
        unit = f'{position} / {kind}' if position else ('VP / Parcel Van' if kind == 'VP' else kind)
        train = re.search(r'\b\d{4,5}\b', asset)
        key = f'parcels|{train.group().lower() if train else asset.lower()}|{unit.lower()}'
        asset_type = kind
    elif book_id == 'catering':
        station = str(at(4) or '').strip()
        category = str(at(6) or '').strip()
        kind = 'MPS' if re.search(r'\bMPS\b', asset, re.I) else asset.upper()
        unit = f'{station} · {category or "Category not stated"}'
        key = f'catering|{asset.lower().replace("-", " ")}|{station.lower()}|{kind.lower()}'
        asset_type = kind
    else:
        category = str(at(2) or '').strip()
        grade = str(at(4) or '').strip()
        asset = ' · '.join(dict.fromkeys(part for part in (grade, source_asset) if part))
        unit = category
        key = f'others|{asset.lower()}|{category.lower()}'
        asset_type = category
    justification = str(at(config[2]) or '').strip()
    identity = f'{book_id}-{sheet}-{row["n"]}'
    if prior and prior['asset'].lower() == asset.lower():
        key, unit, asset_type = prior['key'], prior['unit'], prior['assetType']
    outcome = 'NOT_RECORDED'
    note = justification.lower()
    if re.search(r'no bids?|without bids?|nil bids?', note): outcome = 'NO_BIDS'
    elif re.search(r'terminat', note): outcome = 'TERMINATED'
    elif re.search(r'reject', note): outcome = 'REJECTED'
    elif re.search(r'(?:awarded|sold|accepted)', note): outcome = 'SOLD'
    elif re.search(r'pending|under process', note): outcome = 'PENDING'
    record = {'id': identity, 'key': key,
              'date': parsed.isoformat(), 'dateLabel': date_label,
              'portfolio': book['portfolio'], 'bookId': book_id, 'asset': asset,
              'unit': unit, 'assetType': asset_type,
              'rp': rate, 'previousValue': None, 'justification': justification,
              'outcome': outcome, 'outcomeText': 'Outcome not recorded in source row.',
              'outcomeValue': None, 'sheet': sheet, 'row': row['n'],
              'formula': row['f'].get(str(config[1]))}
    record['search'] = (' '.join(str(value) for value in record.values() if value is not None) + ' ' + row['s']).lower()
    return record


def verify_workbooks(payload: dict, root: Path) -> dict:
    """Compare every populated source cell, formula and hidden-row flag to the extract."""
    expected = {book['id']: book for book in payload['workbooks']}
    if set(expected) != {'parcels', 'catering', 'others'}:
        raise ValueError('Expected the three known Deskflow workbooks.')
    manifest = {}
    total_rows = total_cells = 0
    for book in payload['workbooks']:
        source = (root / book['sourceRelativePath']).resolve()
        if not source.is_relative_to(root.resolve()) or not source.is_file():
            raise ValueError(f"Missing source workbook: {book['sourceRelativePath']}")
        formula_book = load_workbook(source, data_only=False, read_only=False)
        value_book = load_workbook(source, data_only=True, read_only=False)
        try:
            if [s.title for s in formula_book.worksheets] != [s['name'] for s in book['sheets']]:
                raise ValueError(f'Sheet list differs: {source.name}')
            rows_in_book = cells_in_book = 0
            for extracted_sheet in book['sheets']:
                sheet = formula_book[extracted_sheet['name']]
                values_sheet = value_book[sheet.title]
                populated = sorted({cell.row for cell in sheet._cells.values() if cell.value is not None})
                extracted_rows = extracted_sheet['rows']
                if populated != [row['n'] for row in extracted_rows]:
                    raise ValueError(f'Row list differs: {source.name} / {sheet.title}')
                for row in extracted_rows:
                    n = row['n']
                    if bool(sheet.row_dimensions[n].hidden) != row['hidden']:
                        raise ValueError(f'Hidden row differs: {source.name} / {sheet.title} / {n}')
                    actual_values = []
                    actual_formulas = {}
                    for column in range(1, sheet.max_column + 1):
                        cell = sheet._cells.get((n, column))
                        raw = cell.value if cell else None
                        cached = values_sheet._cells.get((n, column))
                        if isinstance(raw, str) and raw.startswith('='):
                            actual_formulas[str(column)] = raw
                            actual_values.append(_serialise(cached.value if cached else None))
                        else:
                            actual_values.append(_serialise(raw))
                        if raw is not None:
                            cells_in_book += 1
                    while actual_values and actual_values[-1] is None:
                        actual_values.pop()
                    if actual_values != row['v'] or actual_formulas != row['f']:
                        raise ValueError(f'Cell differs: {source.name} / {sheet.title} / {n}')
                rows_in_book += len(populated)
            if rows_in_book != book['rowCount'] or cells_in_book != book['cellCount']:
                raise ValueError(f'Count differs: {source.name}')
            total_rows += rows_in_book
            total_cells += cells_in_book
            manifest[book['id']] = {
                'file': book['file'], 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
                'sheets': len(book['sheets']), 'rows': rows_in_book, 'cells': cells_in_book,
            }
        finally:
            formula_book.close()
            value_book.close()
    if (total_rows, total_cells, len(payload['auctionRecords'])) != (
        payload['totalRows'], payload['totalCells'], payload['auctionRecordCount']
    ):
        raise ValueError('Extract totals differ from workbook totals.')
    return manifest
